# -*- coding: utf-8 -*-
"""MIR001 C3b FINAL CLOSURE — test pomoćnik: AI radni predlog → izričito usvajanje → dokument za sud, kroz STVARNE rute
(`POST /api/staging/{predlog}/dokument-za-sud`, `/stavka`, `/procesne-tacke`, `PATCH /tekst`, `/approve`, `/docx`).
Pomoćnik samo poziva rute kao korisnik; ništa ne upisuje u bazu mimo njih."""
import io
import re

from shared import usvajanje_nacrta as UN

SVE_NEPOZNATO = {k: {"odluka": "NEPOZNATO", "tekst": ""} for k in UN.KLJUCEVI_TACAKA}


def bez_oznaka(t: str) -> str:
    """Radnja ADVOKATA u testu: interne oznake spisa (DOK-nn) zamenjuje opisom — sistem ih ne briše sam (R1)."""
    t = re.sub(r"\s*\[\[\s*DOK-\d+\s*\]\]", "", t)
    t = re.sub(r"\s*\((?:DOK-\d+(?:,\s*)?)+\)", "", t)
    return re.sub(r"\bDOK-\d+\b", "dokument spisa", t)


def predlog(k, h, sid) -> dict:
    r = k.get(f"/api/staging/{sid}", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def dokument(k, h, sid_predloga) -> dict:
    r = k.post(f"/api/staging/{sid_predloga}/dokument-za-sud", headers=h)
    assert r.status_code == 200, r.text
    return r.json()


def stavka(k, h, did, sid, odluka="USVOJENO"):
    return k.post(f"/api/staging/{did}/stavka", json={"stavka": sid, "odluka": odluka}, headers=h)


def tacke(k, h, did, odluke=None):
    return k.post(f"/api/staging/{did}/procesne-tacke", json={"tacke": odluke or SVE_NEPOZNATO}, headers=h)


def sacuvan(k, h, did) -> dict:
    return predlog(k, h, did)["usvajanje"]


def odobri(k, h, did, sha=None, rev=None):
    d = sacuvan(k, h, did)
    return k.post(f"/api/staging/{did}/approve", json={"sha256": sha or d["sha256"], "rev": d["rev"] if rev is None else rev},
                  headers=h)


def docx_pasusi(sadrzaj: bytes) -> list:
    import docx
    return [p.text for p in docx.Document(io.BytesIO(sadrzaj)).paragraphs]


def docx_podnozje(sadrzaj: bytes) -> str:
    import docx
    d = docx.Document(io.BytesIO(sadrzaj))
    return "\n".join(p.text for s in d.sections for p in s.footer.paragraphs)


def ceo_tok(k, h, sid_predloga, usvoji, izmeni=None, odluke=None):
    """Usvoji navedene stavke → (izmena teksta) → procesne tačke → odobri → DOCX. Vraća (dokument, docx odgovor)."""
    d = dokument(k, h, sid_predloga)
    for s in usvoji:
        r = stavka(k, h, d["id"], s)
        assert r.status_code == 200, r.text
    if izmeni:
        t = izmeni(sacuvan(k, h, d["id"])["tekst"])
        r = k.patch(f"/api/staging/{d['id']}/tekst", json={"tekst": t}, headers=h)
        assert r.status_code == 200, r.text
    r = tacke(k, h, d["id"], odluke)
    assert r.status_code == 200, r.text
    r = odobri(k, h, d["id"])
    assert r.status_code == 200 and r.json()["status"] == "approved", r.text
    return sacuvan(k, h, d["id"]), k.post(f"/api/staging/{d['id']}/docx", headers=h)
