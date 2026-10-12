# -*- coding: utf-8 -*-
"""VCR002 Task 02 — kontrolisan sintetički korpus kroz STVARNE Smart Intake rute, STVARAN intake worker
i (gde je potreban) STVARAN Tesseract. Dopunjuje tests/test_ns005_t14_smart_intake.py samo onim što tamo
nije dokazano: DOCX, mešovit PDF (digitalna + skenirana ćirilična strana) do teksta prikačenog predmetu,
dva fajla istog imena a različitog sadržaja, prevelik fajl, oštećen PDF, pad upisa POSLE ekstrakcije
(prolazna greška baze → ponovni pokušaj, bez duplikata), pad upisa pri prikačivanju (nikad lažan uspeh),
ponovljeno prikačivanje bez ključa i tuđ predmet.

Sva imena, brojevi predmeta i činjenice su izmišljeni. Bez Tesseract-a se preskaču SAMO OCR testovi;
sa VINDEX_REQUIRE_OCR=1 (produkciona slika u CI) nedostatak je pad.
"""
import io
import os
import unicodedata

import pytest

from tests.ns005_harness import zaglavlje, SPOLJNI_POKUSAJI
from tests.test_ns005_t14_smart_intake import (  # noqa: F401 — `ok` je pytest fixture
    PA, PB, _OCR, _obradi, _otpremi, _skeniran_pdf, ok,
)

treba_ocr = pytest.mark.skipif(_OCR is None, reason="Tesseract nije dostupan — OCR deo korpusa preskočen")

DIGITALNI = [
    "PRIVREDNI SUD U KRAGUJEVCU",
    "Broj predmeta: P 4711/2026",
    "Tuzilac: Zlatibor Fiktivic iz Kraljeva",
    "Tuzeni: Primer Gradnja DOO iz Cacka",
    "Rok za odgovor na tuzbu je 30 dana od dana dostavljanja.",
]
SKENIRANI_CIR = [
    "ЗАПИСНИК СА РОЧИШТА",
    "Број предмета: П 4711/2026",
    "Сведок Драгослав Измишљени изјавио је",
    "да је роба испоручена 3. марта 2026. године.",
    "Рочиште је одложено за 20. децембар 2026. године.",
]
RECI_CIR = ["записник", "рочишта", "предмета", "сведок", "драгослав", "изјавио", "роба", "испоручена",
            "марта", "године", "одложено", "децембар"]
DOCX_PASUS = ("Ugovor o zakupu poslovnog prostora broj 17/2026 zakljucen je izmedju fiktivnog zakupodavca "
              "Milorada Primerovica i zakupca Test Trgovina DOO, sa mesecnom zakupninom od 120.000 dinara.")


def _norm(t):
    t = "".join(c for c in unicodedata.normalize("NFKD", t or "") if not unicodedata.combining(c))
    return " ".join(t.lower().split())


def _digitalna_strana(doc, redovi):
    import fitz
    s = doc.new_page(width=595, height=842)
    for i, red in enumerate(redovi):
        s.insert_text((50, 80 + i * 30), red, fontsize=14)


def _digitalni_pdf(redovi=DIGITALNI):
    import fitz
    d = fitz.open()
    _digitalna_strana(d, redovi)
    b = io.BytesIO()
    d.save(b)
    d.close()
    return b.getvalue()


def _mesoviti_pdf():
    """Strana 1: pravi tekstualni sloj (latinica). Strana 2: rasterizovan ćirilični sken bez sloja."""
    import fitz
    sken = fitz.open(stream=_skeniran_pdf(SKENIRANI_CIR), filetype="pdf")
    d = fitz.open()
    _digitalna_strana(d, DIGITALNI)
    d.insert_pdf(sken)
    sken.close()
    b = io.BytesIO()
    d.save(b)
    d.close()
    return b.getvalue()


def _docx(pasus=DOCX_PASUS):
    import docx
    d = docx.Document()
    d.add_paragraph("UGOVOR O ZAKUPU")
    d.add_paragraph(pasus)
    b = io.BytesIO()
    d.save(b)
    return b.getvalue()


def _job(r):
    assert r.status_code == 202, r.text
    x = r.json()["rezultati"][0]
    assert x["ok"] is True, x
    return x["job_id"]


def _prikaci(k, jid, predmet=PA, korisnik="A"):
    if k.get(f"/api/smart-intake/jobs/{jid}", headers=zaglavlje(korisnik)).json()["job"]["status"] == "awaiting_review":
        assert k.post(f"/api/smart-intake/jobs/{jid}/review/resolve", headers=zaglavlje(korisnik)).status_code == 200
    return k.post(f"/api/smart-intake/jobs/{jid}/finalize", json={"predmet_id": predmet}, headers=zaglavlje(korisnik))


def _dokumenti(b, predmet=PA):
    return [d for d in b.tabele["predmet_dokumenti"] if d["predmet_id"] == predmet]


# ── A: digitalni PDF — tačan tekst, OCR se ne koristi ───────────────────────────────────────────────

def test_A_digitalni_pdf_tacan_tekst_bez_ocr(ok):
    k, b = ok
    jid = _job(_otpremi(k, "A", _digitalni_pdf(), ime="tuzba digitalna.pdf"))
    assert _obradi() is True
    d = k.get(f"/api/smart-intake/jobs/{jid}", headers=zaglavlje("A")).json()
    assert d["dokument"]["ocr_koriscen"] is False
    assert _prikaci(k, jid).status_code == 200
    [dok] = _dokumenti(b)
    tekst = _norm(dok.get("tekst_sadrzaj"))
    for red in DIGITALNI:
        assert _norm(red) in tekst, red


# ── D: DOCX — poznat pasus doslovno u tekstu dokumenta predmeta ─────────────────────────────────────

def test_D_docx_poznat_pasus_do_teksta_predmeta(ok):
    k, b = ok
    jid = _job(_otpremi(k, "A", _docx(), ime="ugovor.docx",
                        tip="application/vnd.openxmlformats-officedocument.wordprocessingml.document"))
    assert _obradi() is True
    assert b.tabele["intake_jobs"][0]["status"] in ("awaiting_review", "completed")
    assert _prikaci(k, jid).status_code == 200
    [dok] = _dokumenti(b)
    assert _norm(DOCX_PASUS) in _norm(dok.get("tekst_sadrzaj")) and dok["naziv_fajla"] == "ugovor.docx"


# ── C: mešovit PDF — skenirana ćirilična strana NIJE tiho izgubljena ────────────────────────────────

@treba_ocr
def test_C_mesoviti_pdf_obe_strane_u_tekstu_predmeta_i_posle_ponovnog_citanja(ok):
    k, b = ok
    from uploaded_doc import extractor
    jezik = extractor._detect_ocr_lang()
    jid = _job(_otpremi(k, "A", _mesoviti_pdf(), ime="spis mesoviti.pdf"))
    assert _obradi() is True
    d = k.get(f"/api/smart-intake/jobs/{jid}", headers=zaglavlje("A")).json()
    assert d["dokument"]["ocr_koriscen"] is True          # skenirana strana prošla kroz OCR
    assert _prikaci(k, jid).status_code == 200
    [dok] = _dokumenti(b)
    tekst = (dok.get("tekst_sadrzaj") or "").lower()
    assert _norm(DIGITALNI[1]) in _norm(tekst)            # digitalna strana doslovno
    if "srp" in jezik.split("+"):
        udeo = sum(1 for r in RECI_CIR if r in tekst) / len(RECI_CIR)
        assert udeo >= 0.75, (jezik, udeo, tekst[-300:])  # OCR tolerancija, ne bajt-identičnost
    elif os.getenv("VINDEX_EXPECT_SRP") == "1":
        pytest.fail(f"produkciona slika bez ćiriličnog paketa: {jezik!r}")
    else:
        assert len(tekst) > len(" ".join(DIGITALNI)) + 40, "skenirana strana je tiho izostavljena"
    # „reload“: ponovno čitanje istog posla daje isto stanje, ne novo
    d2 = k.get(f"/api/smart-intake/jobs/{jid}", headers=zaglavlje("A")).json()
    assert d2["job"]["status"] == d["job"]["status"] or d2["job"]["status"] == "completed"
    assert len(_dokumenti(b)) == 1


# ── Isto ime, različit sadržaj: dva dokumenta, ne jedan ─────────────────────────────────────────────

def test_isto_ime_razlicit_sadrzaj_dva_razlicita_dokumenta(ok):
    k, b = ok
    j1 = _job(_otpremi(k, "A", _digitalni_pdf(), ime="tuzba.pdf"))
    j2 = _job(_otpremi(k, "A", _digitalni_pdf(DIGITALNI[:3] + ["Drugi dokument istog imena, drugaciji sadrzaj."]), ime="tuzba.pdf"))
    assert j1 != j2
    while _obradi():
        pass
    assert _prikaci(k, j1).status_code == 200 and _prikaci(k, j2).status_code == 200
    docs = _dokumenti(b)
    assert len(docs) == 2 and len({d["id"] for d in docs}) == 2
    assert sum("drugaciji sadrzaj" in _norm(d.get("tekst_sadrzaj")) for d in docs) == 1


# ── Granice prijema ─────────────────────────────────────────────────────────────────────────────────

def test_prevelik_fajl_odbijen_bez_posla_i_bez_skladista(ok):
    k, b = ok
    r = _otpremi(k, "A", b"%PDF-1.4\n" + b"0" * (25 * 1024 * 1024 + 1), ime="ogroman.pdf")
    assert r.status_code == 202 and r.json()["rezultati"][0]["ok"] is False
    assert "prevelik" in r.json()["rezultati"][0]["greska"].lower()
    assert b.tabele["intake_jobs"] == [] and b.storage.sadrzaj == {}


def test_osteceni_pdf_nikad_lazan_uspeh(ok):
    k, b = ok
    jid = _job(_otpremi(k, "A", b"%PDF-1.4\n%ostecen\x00\x01 nije pdf", ime="osteceni.pdf"))
    _obradi()
    posao = b.tabele["intake_jobs"][0]
    assert posao["status"] != "completed", posao
    d = k.get(f"/api/smart-intake/jobs/{jid}", headers=zaglavlje("A")).json()
    assert d["entiteti"] == []
    if posao["status"] == "awaiting_review":
        assert d["potrebna_provera"]["razlog"] in ("ocr_failed", "extraction_failed", "low_confidence"), d["potrebna_provera"]
    assert _dokumenti(b) == []


# ── Pad upisa baze posle ekstrakcije: ponovni pokušaj, bez duplikata ────────────────────────────────

def test_pad_upisa_posle_ekstrakcije_ponovni_pokusaj_bez_duplikata(ok):
    k, b = ok
    jid = _job(_otpremi(k, "A", _digitalni_pdf(), ime="tuzba.pdf"))
    b.greske["intake_documents"] = Exception('{"code": "57014", "message": "canceling statement due to statement timeout"}')
    assert _obradi() is True
    posao = b.tabele["intake_jobs"][0]
    assert posao["status"] == "received" and posao["attempts"] == 1 and posao["last_error"]   # zakazan ponovni pokušaj
    del b.greske["intake_documents"]
    assert k.get(f"/api/smart-intake/jobs/{jid}", headers=zaglavlje("A")).json()["job"]["status"] not in ("completed", "awaiting_review")
    posao["next_retry_at"] = None                                                              # backoff istekao
    assert _obradi() is True
    assert posao["status"] in ("awaiting_review", "completed")
    assert len([d for d in b.tabele["intake_documents"] if d.get("intake_job_id") == jid]) == 1
    tipovi = [e["entity_type"] for e in b.tabele["extracted_entities"]]
    assert len(tipovi) == len(set(tipovi)), tipovi                                            # entiteti nisu udvostručeni


# ── Pad upisa pri prikačivanju: nikad „povezano“ ako dokument nije sačuvan ──────────────────────────

def test_pad_upisa_pri_prikacivanju_nije_uspeh_a_ponovni_pokusaj_pravi_jedan_dokument(ok):
    k, b = ok
    jid = _job(_otpremi(k, "A", _digitalni_pdf(), ime="tuzba.pdf"))
    assert _obradi() is True
    b.greske["predmet_dokumenti"] = Exception('{"code": "57014", "message": "canceling statement due to statement timeout"}')
    f = _prikaci(k, jid)
    del b.greske["predmet_dokumenti"]
    assert _dokumenti(b) == []
    if f.status_code == 200:
        assert f.json().get("dokumenata_povezano", 0) == 0, f.json()   # ne sme tvrditi povezivanje
    assert "57014" not in f.text and "statement timeout" not in f.text  # bez sirove greške baze
    # Zauzeće finalizacije (migracija 092) se NE oslobađa posle neočekivane greške; samo-isceljuje se posle
    # prozora zastarelosti (120 s). Do tada je odgovor iskren 409 „u toku“, nikad drugi dokument.
    rano = _prikaci(k, jid)
    assert rano.status_code == 409 and _dokumenti(b) == []
    from datetime import datetime, timedelta, timezone
    b.tabele["intake_jobs"][0]["finalizing_at"] = (datetime.now(timezone.utc) - timedelta(seconds=121)).isoformat()
    f2 = _prikaci(k, jid)
    assert f2.status_code == 200, f2.text
    assert f2.json()["dokumenata_povezano"] == 1 and len(_dokumenti(b)) == 1


def test_ponovljeno_prikacivanje_bez_kljuca_ne_pravi_drugi_dokument(ok):
    k, b = ok
    jid = _job(_otpremi(k, "A", _digitalni_pdf(), ime="tuzba.pdf"))
    assert _obradi() is True
    assert _prikaci(k, jid).status_code == 200
    f2 = k.post(f"/api/smart-intake/jobs/{jid}/finalize", json={"predmet_id": PA}, headers=zaglavlje("A"))
    assert f2.status_code in (200, 409), f2.text
    assert len(_dokumenti(b)) == 1 and len(b.vektori) == 1


def test_prikacivanje_u_tudj_predmet_odbijeno_i_bez_traga(ok):
    k, b = ok
    jid = _job(_otpremi(k, "A", _digitalni_pdf(), ime="tuzba.pdf"))
    assert _obradi() is True
    f = _prikaci(k, jid, predmet=PB)
    assert f.status_code in (403, 404), f.text
    assert b.tabele["predmet_dokumenti"] == [] and b.vektori == []


def test_nema_spoljne_mreze():
    assert SPOLJNI_POKUSAJI == []
