# -*- coding: utf-8 -*-
"""MIR001 C3b FINAL CLOSURE TASK-03 — merenje granice prenosa (PREREGISTRACIJA.md u istom folderu).

Režimi (`FC_MOD`):
  replay  — deo A: 100 SIROVIH izlaza gpt-4o iz GATE-03 kroz NOVU rutu (model zamenjen snimljenim izlazom), mehanički
            advokat (pravilo iz PREREGISTRACIJA §2), odobrenje, DOCX; plus A2 (ubačena netačna tvrdnja, indeks 00).
  gen     — deo B: po jedan STVARAN poziv gpt-4o po scenariju kroz novu rutu; čuva sirov izlaz i ono što advokat vidi.
  primeni — deo B1: odluke nezavisnog „advokata“ (`b1/odluke.json`) se primenjuju kroz stvarne rute nad ponovljenim
            snimljenim izlazom; odobrenje, DOCX.
Sve ostalo je stvarno (rute, procesno stanje, staging, DOCX); zamenjeni su baza u memoriji, pravna baza, ocena kvaliteta
i naplata (kao GATE-03). Ključ se čita iz `.env` i ne ispisuje.

Pokretanje iz korena repozitorijuma (OCR okruženje uključeno):
  FC_MOD=replay python -X utf8 -m pytest -p tests.conftest -p no:randomly -p no:cacheprovider -q \
      docs/v2-execution/mir001_c3b_model/final_closure/t03/t03_fc.py
"""
import asyncio
import importlib.util
import io
import json
import os
import re
import types

import pytest

from tests import mir001_golden_case as gc
from tests.test_mir001_t02_prijem_spisa import primi_spis
from tests.test_osb002_solo_tok import A, solo  # noqa: F401

OVDE = os.path.dirname(os.path.abspath(__file__))
STARI = os.path.join(OVDE, "..", "..", "strukturalno", "t03")
_spec = importlib.util.spec_from_file_location("gen_t03_gate03", os.path.join(STARI, "generisi_t03.py"))
G = importlib.util.module_from_spec(_spec)
_spec.loader.exec_module(G)                      # spis(sc), ZAHTEV, SCENARIJI, _kljuc_iz_env — ISTO kao GATE-03

MOD = os.environ.get("FC_MOD", "replay")
MATERIJALNA = ("M1", "M2", "M3", "M4", "M5", "M7")
UBACENA = "Tuzilac je dana 20.07.2026. godine uplatio predujam od 60.000 dinara, a sud je predlog za dopunu vestacenja odbio."
SVE_NEPOZNATO = None                             # popunjava se iz shared.usvajanje_nacrta


def _n(t):
    return " ".join(str(t or "").replace("„", "").replace("“", "").replace('"', "").split()).lower()


def _citati_gate03():
    """{izlaz: [normalizovani citati]} — materijalni nalazi BILO KOG slepog označivača GATE-03."""
    d = json.load(open(os.path.join(STARI, "spojeno_t03.json"), encoding="utf-8"))
    out = {}
    for x in d["nalazi"]:
        if x["merilo"] not in MATERIJALNA:
            continue
        c = x.get("citat") or ""
        for deo in (c if isinstance(c, list) else str(c).split(" || ")):
            deo = _n(deo).strip(". ")
            if len(deo) >= 12:
                out.setdefault(x["izlaz"], []).append(deo)
    return out


@pytest.fixture
def okruzenje(solo, monkeypatch):
    k, b = solo
    import app.services.retrieve as R
    import routers.case_dna as cd
    import routers.evidence as ev
    import routers.multi_agent as MA
    import services.quality_gate as QG
    import shared.usage as US
    from tests.test_mir001_t03_kontinuitet import _klasifikuj
    from tests.test_mir001_t04_savetnik_spis import _genome_bez_narativa
    monkeypatch.setattr(cd, "_pozovi_genome_api", _genome_bez_narativa(b.genome_pozivi))
    monkeypatch.setattr(ev, "_klasifikuj_dokument", _klasifikuj)

    async def _politika(feature):
        return {"aktivno": True, "status": "ACTIVE", "feature_type": "CORE", "min_plan": "free", "krediti": 1}

    async def _stanje(uid, email):
        return 100
    naplate, pozivi = [], []

    async def _consume(*a, **kw):
        naplate.append(1)

    async def _kvalitet(tekst, tip):
        return {"confidence_score": 0.5, "detail": {}}
    monkeypatch.setattr(US, "get_policy", _politika)
    monkeypatch.setattr(US.UsageService, "balance", staticmethod(_stanje))
    monkeypatch.setattr(US.UsageService, "consume", staticmethod(_consume))
    monkeypatch.setattr(R, "retrieve_documents", lambda *a, **kw: ([], {}))
    monkeypatch.setattr(QG, "evaluate_draft_quality", _kvalitet)
    pravi = MA._pozovi_nacrt_api
    b.snimljen, b.sirovo, b.greske_modela = None, None, []

    def _model(client, system, user_msg):
        pozivi.append(user_msg)
        if b.snimljen is not None:               # replay / primeni: snimljen SIROV izlaz modela
            return types.SimpleNamespace(choices=[types.SimpleNamespace(message=types.SimpleNamespace(content=b.snimljen[0]),
                                                                        finish_reason=b.snimljen[1])])
        import socket
        from tests import ns005_harness as H
        cuvar = socket.socket.connect
        try:                                     # stvaran poziv: mreža otvorena samo za njegovo trajanje (kao GATE-03)
            socket.socket.connect = H._orig_connect
            resp = pravi(client, system, user_msg)
            b.sirovo = [resp.choices[0].message.content, resp.choices[0].finish_reason]
            return resp
        except BaseException as e:
            b.greske_modela.append(f"{type(e).__name__} {getattr(e, 'status_code', '')}".strip())
            raise
        finally:
            socket.socket.connect = cuvar
    monkeypatch.setattr(MA, "_pozovi_nacrt_api", _model)
    if MOD == "gen":
        monkeypatch.setenv("OPENAI_API_KEY", G._kljuc_iz_env())
    b.odlozeni = []

    class _Asyncio:
        def __getattr__(self, ime):
            return getattr(asyncio, ime)

        def create_task(self, coro):
            b.odlozeni.append(coro)
    monkeypatch.setattr(MA, "asyncio", _Asyncio())
    import routers.drafting as D
    monkeypatch.setattr(D, "asyncio", _Asyncio())     # audit zapisi usvajanja se izvršavaju eksplicitno
    b.podrazumevano.setdefault("staging_memory", {}).update({"status": "pending", "is_lawyer_approved": False,
                                                            "pinecone_indexed": False, "quality_detail": {}, "naziv": ""})
    return k, b, naplate, pozivi


def _izvrsi(b):
    while b.odlozeni:
        asyncio.run(b.odlozeni.pop(0))


def _predlog(k, b, sc, snimljen=None):
    """Spis scenarija → `/api/agents/run` → AI radni predlog na pregledu. Vraća (odgovor rute, red predloga, prikaz)."""
    pid, doks, _ = primi_spis(k, b, dokumenti=G.spis(sc))
    b.snimljen = snimljen
    r = k.post("/api/agents/run", json={"agent": "drafting", "task": G.ZAHTEV, "predmet_id": pid}, headers=A)
    _izvrsi(b)
    st = [s for s in b.tabele["staging_memory"] if s.get("predmet_id") == pid and s.get("tip") == "nacrt_iz_spisa"]
    prikaz = k.get(f"/api/staging/{st[0]['id']}", headers=A).json() if st else None
    return r, (st[0] if st else None), prikaz


def _docx(sadrzaj):
    import docx
    d = docx.Document(io.BytesIO(sadrzaj))
    return ([p.text for p in d.paragraphs], [p.text for s in d.sections for p in s.footer.paragraphs],
            [p.text for s in d.sections for p in s.header.paragraphs])


def _ocekivani_pasusi(tekst):
    out = []
    for red in tekst.split("\n"):
        for pre in ("### ", "## ", "# "):
            if red.startswith(pre):
                red = red[len(pre):]
                break
        red = red.replace("**", "")
        if red.strip():
            out.append(red)
    return out


def _usvoji_i_izvezi(k, b, prikaz, odluke_stavki, tacke, zamene=(), dodatak=""):
    """Kroz STVARNE rute: dokument za sud → odluke o stavkama → (izmena advokata) → tačke → odobrenje → DOCX.
    Vraća zapis sa svim proverama K."""
    from shared import usvajanje_nacrta as UN
    from shared.nacrt_za_sud import GRANICA
    sid = prikaz["id"]
    z = {"provere": {}, "odluke": odluke_stavki}
    pr_red = next(s for s in b.tabele["staging_memory"] if s["id"] == sid)
    sha_predloga = UN.sha(pr_red["tekst"])
    # K4/K5: predlog se ne odobrava, ne menja, ne izvozi; dokument za sud je prazan dok se ne usvoji
    z["provere"]["K4_predlog_zakljucan"] = (k.post(f"/api/staging/{sid}/approve", json={}, headers=A).status_code == 409
                                           and k.patch(f"/api/staging/{sid}/tekst", json={"tekst": "x"}, headers=A).status_code == 409
                                           and k.post(f"/api/staging/{sid}/docx", headers=A).status_code == 409
                                           and k.post("/api/nacrti/export/docx", json={"tekst": pr_red["tekst"][:5000], "naslov": "X",
                                                                                       "tip": "nacrt_iz_spisa"}, headers=A).status_code == 409)
    d = k.post(f"/api/staging/{sid}/dokument-za-sud", headers=A).json()
    z["provere"]["K5_prazan_na_pocetku"] = d.get("tekst") == "" and d.get("procesne_tacke") == {}
    z["provere"]["K5_bez_tacaka_nema_odobrenja"] = k.post(f"/api/staging/{d['id']}/approve",
                                                          json={"sha256": d["sha256"], "rev": d["rev"]}, headers=A).status_code == 409
    for s in prikaz["predlog"]["stavke"]:
        o = odluke_stavki.get(s["id"])
        if o in ("USVOJENO", "ODBACENO"):
            r = k.post(f"/api/staging/{d['id']}/stavka", json={"stavka": s["id"], "odluka": o}, headers=A)
            assert r.status_code == 200, r.text
    tekst = k.get(f"/api/staging/{d['id']}", headers=A).json()["usvajanje"]["tekst"]
    z["zamene_nisu_nadjene"] = [staro for staro, _ in zamene if not re.search(staro, tekst)]
    izmenjen = _primeni_zamene(tekst, zamene)
    if dodatak:
        izmenjen = (izmenjen.rstrip() + "\n\n" + dodatak) if izmenjen.strip() else dodatak
    if izmenjen != tekst and izmenjen.strip():
        assert k.patch(f"/api/staging/{d['id']}/tekst", json={"tekst": izmenjen}, headers=A).status_code == 200
    r = k.post(f"/api/staging/{d['id']}/procesne-tacke", json={"tacke": tacke}, headers=A)
    z["tacke_status"] = r.status_code
    u = k.get(f"/api/staging/{d['id']}", headers=A).json()["usvajanje"]
    z["razlozi_protiv_odobrenja"] = u.get("razlozi_protiv_odobrenja")
    ra = k.post(f"/api/staging/{d['id']}/approve", json={"sha256": u["sha256"], "rev": u["rev"]}, headers=A)
    z["odobrenje"] = ra.status_code
    _izvrsi(b)
    red = next(s for s in b.tabele["staging_memory"] if s["id"] == d["id"])
    z["odobren_tekst"] = red["tekst"]
    z["status"] = red["status"]
    if ra.status_code != 200:
        return z
    dx = k.post(f"/api/staging/{d['id']}/docx", headers=A)
    z["docx_status"] = dx.status_code
    pasusi, podnozje, zaglavlje = _docx(dx.content) if dx.status_code == 200 else ([], [], [])
    z["docx_pasusi"] = pasusi
    sve = _n("\n".join(pasusi))
    odobren = _n(red["tekst"])
    # K1: neusvojena stavka (≥ 30 znakova) nije ni u odobrenom tekstu ni u DOCX-u
    neusvojene = [s for s in prikaz["predlog"]["stavke"] if odluke_stavki.get(s["id"]) != "USVOJENO" and len(_n(_zd(s))) >= 30]
    z["provere"]["K1_neusvojeno_nije_u_dokumentu"] = all(all(_n(x) not in odobren and _n(x) not in sve for x in (s["tekst"], _zd(s)))
                                                         for s in neusvojene)
    # K2: DOCX = odobren tekst, bez sistemskog teksta
    z["provere"]["K2_docx_jednak_odobrenom"] = [p for p in pasusi if p.strip()] == _ocekivani_pasusi(red["tekst"])
    z["provere"]["K2_bez_sistemskog_teksta"] = (not any(x in "\n".join(pasusi) for x in ("NACRT IZ SPISA", GRANICA[:20], "[IZOSTAVLJENO", "[[PS-",
                                                                                        "INTERNE NAPOMENE", "Dokument za sud —"))
                                              and not re.search(r"\bDOK-\d", "\n".join(pasusi))
                                              and "".join(podnozje).strip() == "" and "Vindex" not in "".join(podnozje + zaglavlje))
    # K3: svaka usvojena stavka je u odobrenom tekstu (posle izmene advokata)
    usvojene = [s for s in prikaz["predlog"]["stavke"] if odluke_stavki.get(s["id"]) == "USVOJENO"]
    pr = izmenjen if izmenjen.strip() else tekst
    # K1b (GATE-03 N1): odobren tekst je TAČNO spoj usvojenih stavki (redosled usvajanja, „\n\n“) uz advokatove zamene i
    # dodatak — i delimično procurela neusvojena stavka bi ovde pala
    rekon = _primeni_zamene("\n\n".join(_zd(s) for s in usvojene), zamene)
    if dodatak:
        rekon = (rekon.rstrip() + "\n\n" + dodatak) if rekon.strip() else dodatak
    z["provere"]["K1b_odobren_tacno_iz_usvojenog"] = red["tekst"] == rekon
    z["provere"]["K3_usvojeno_nije_izgubljeno"] = all(_n(_primeni_zamene(_zd(s), zamene)) in _n(red["tekst"]) for s in usvojene) and red["tekst"] == pr
    # K4: predlog nepromenjen; K7: bez VERIFIED, odobrenje zapisano
    z["provere"]["K4_predlog_nepromenjen"] = UN.sha(next(s for s in b.tabele["staging_memory"] if s["id"] == sid)["tekst"]) == sha_predloga
    qd = json.dumps(red.get("quality_detail") or {}, ensure_ascii=False)
    z["provere"]["K7_bez_verified"] = "VERIFIED" not in qd and (red["quality_detail"]["usvajanje"]["odobreno"] or {}).get("sha256") == UN.sha(red["tekst"])
    return z


def _zd(s):
    """Tekst koji usvajanje unosi u dokument (R1: `tekst_za_dokument`; stariji snimak: `tekst`)."""
    return s.get("tekst_za_dokument") or s["tekst"]


# R1: mehanički advokat zamenjuje interne oznake spisa (sistem ih ne briše; bez zamene odobrenje ne prolazi)
ZAMENE_OZNAKA = ((r"\s*\[\[\s*DOK-\d+\s*\]\]", ""), (r"\s*\((?:DOK-\d+(?:,\s*)?)+\)", ""), (r"\bDOK-\d+\b", "dokument spisa"))


def _primeni_zamene(t, zamene):
    for staro, novo in zamene:
        t = re.sub(staro, novo, t)
    return t


def _k4_ceo_predlog(sirov, prikaz):
    """Svaki deo SIROVOG izlaza modela (oznake [[PS-nn]] izuzete — server ih zamenjuje izvodom) postoji u predlogu."""
    try:
        tekst = json.loads(sirov)["tekst_za_sud"]
    except Exception:
        return None
    predlog = _n("\n\n".join(s["tekst"] for s in prikaz["predlog"]["stavke"]))
    fale = []
    for red in tekst.split("\n"):
        for deo in re.split(r"\[\[\s*PS-\d{1,3}\s*\]\]", red):
            if len(_n(deo).strip(" .,;:()")) >= 3 and _n(deo).strip(" .,;:()") not in predlog:
                fale.append(deo.strip()[:80])
    return fale


def _sve_nepoznato():
    from shared import usvajanje_nacrta as UN
    return {kk: {"odluka": "NEPOZNATO", "tekst": ""} for kk in UN.KLJUCEVI_TACAKA}


# ── Deo A: replay 100 + A2 ───────────────────────────────────────────────────────────────────────────────────────
STARI_IZLAZI = sorted(f[:-5] for f in os.listdir(os.path.join(STARI, "izlazi")) if f.endswith(".json") and not f.startswith("SUVO_"))


@pytest.mark.skipif(MOD != "replay", reason="FC_MOD")
@pytest.mark.parametrize("ime", STARI_IZLAZI + [f"A2_{sc}" for sc in G.SCENARIJI])
def test_replay(okruzenje, ime):
    k, b, naplate, pozivi = okruzenje
    a2 = ime.startswith("A2_")
    sc = ime[3:] if a2 else ime.split("_")[0]
    izvor = f"{sc}_00" if a2 else ime
    staro = json.load(open(os.path.join(STARI, "izlazi", izvor + ".json"), encoding="utf-8"))
    sirov = list(staro["sirov_izlaz_modela"])
    if a2:
        t = json.loads(sirov[0])
        t["tekst_za_sud"] = t["tekst_za_sud"].rstrip() + "\n\n" + UBACENA
        sirov[0] = json.dumps(t, ensure_ascii=False)
    r, pr_red, prikaz = _predlog(k, b, sc, snimljen=sirov)
    z = {"izlaz": ime, "scenarij": sc, "ruta": r.status_code, "pozivi_modela": len(pozivi), "naplate": len(naplate)}
    if r.status_code != 200 or not prikaz:
        z["napomena"] = "predlog nije napravljen"
        _upisi("replay", ime, z)
        return
    z["status_teksta"] = r.json().get("status_teksta")
    z["K4_fali_iz_sirovog"] = _k4_ceo_predlog(sirov[0], prikaz)
    citati = _citati_gate03().get(izvor, []) if not a2 else _citati_gate03().get(izvor, [])
    odluke = {}
    for s in prikaz["predlog"]["stavke"]:
        nt = _n(s["tekst"])
        odbaci = ("⚠" in s["tekst"] or "NE POSTOJI]]" in s["tekst"] or any(c in nt for c in citati)
                  or (a2 and _n(UBACENA) in nt))
        odluke[s["id"]] = "ODBACENO" if odbaci else "USVOJENO"
    if a2:
        ub = [s for s in prikaz["predlog"]["stavke"] if _n(UBACENA) in _n(s["tekst"])]
        z["A2_ubacena_je_ai_predlog"] = len(ub) == 1 and ub[0]["status"] == "AI_RADNI_PREDLOG"
    zz = _usvoji_i_izvezi(k, b, prikaz, odluke, _sve_nepoznato(), zamene=((r"\[PROVERITI[^\]]*\]", "(upisao advokat)"),) + ZAMENE_OZNAKA)
    z.update(zz)
    z["odbaceno"] = sum(1 for v in odluke.values() if v == "ODBACENO")
    z["usvojeno"] = sum(1 for v in odluke.values() if v == "USVOJENO")
    z["pozivi_modela_posle"], z["naplate_posle"] = len(pozivi), len(naplate)
    if a2:
        z["A2_ubacena_nije_u_dokumentu"] = _n(UBACENA) not in _n(z.get("odobren_tekst", "")) and _n(UBACENA) not in _n("\n".join(z.get("docx_pasusi", [])))
    _upisi("replay", ime, z)


# ── Deo B: stvaran uzorak ────────────────────────────────────────────────────────────────────────────────────────
@pytest.mark.skipif(MOD != "gen", reason="FC_MOD")
@pytest.mark.parametrize("sc", G.SCENARIJI)
def test_gen(okruzenje, sc):
    k, b, naplate, pozivi = okruzenje
    from shared.nacrt_za_sud import GRANICA
    r, pr_red, prikaz = _predlog(k, b, sc, snimljen=None)
    z = {"scenarij": sc, "ruta": r.status_code, "pozivi_modela": len(pozivi), "naplate": len(naplate),
         "greske_modela": b.greske_modela, "sirov_izlaz_modela": b.sirovo,
         "dokumenti_spisa": [{"oznaka": x["oznaka"], "ime": x["ime"], "tekst": gc.tekst(x)} for x in G.spis(sc)]}
    if prikaz:
        z["predlog"] = prikaz["predlog"]
        z["interno"] = GRANICA + pr_red["tekst"].split(GRANICA, 1)[1] if GRANICA in pr_red["tekst"] else ""
        z["status_teksta"] = prikaz.get("status_teksta")
        z["K4_fali_iz_sirovog"] = _k4_ceo_predlog(b.sirovo[0], prikaz) if b.sirovo else None
    _upisi("gen", sc, z)
    assert r.status_code in (200, 502), r.text[:300]


@pytest.mark.skipif(MOD != "primeni", reason="FC_MOD")
@pytest.mark.parametrize("sc", G.SCENARIJI)
def test_primeni(okruzenje, sc):
    k, b, naplate, pozivi = okruzenje
    gen = json.load(open(os.path.join(OVDE, "gen", sc + ".json"), encoding="utf-8"))
    odl = json.load(open(os.path.join(OVDE, "b1", "odluke.json"), encoding="utf-8"))[sc]
    r, pr_red, prikaz = _predlog(k, b, sc, snimljen=gen["sirov_izlaz_modela"])
    assert prikaz["predlog"]["sha256"] == gen["predlog"]["sha256"]          # isti predlog koji je advokat video
    odluke = {sid: ("USVOJENO" if v == "USVAJAM" else "ODBACENO") for sid, v in odl["stavke"].items()}
    tacke = {kk: {"odluka": v["odluka"], "tekst": v.get("tekst", "")} for kk, v in odl["tacke"].items()}
    zamene = [(re.escape(x["staro"]), x["novo"].replace("\\", "\\\\")) for x in odl.get("zamene") or [] if x.get("staro")] + list(ZAMENE_OZNAKA)
    z = _usvoji_i_izvezi(k, b, prikaz, odluke, tacke, zamene=zamene, dodatak=odl.get("dodatni_tekst") or "")
    z["scenarij"], z["pozivi_modela"], z["naplate"] = sc, len(pozivi), len(naplate)
    _upisi("primeni", sc, z)


def _upisi(mod, ime, z):
    d = os.path.join(OVDE, os.environ.get("FC_IZLAZ") or mod)     # FC_IZLAZ: provera harness-a mutacijama (ne gazi merenje)
    os.makedirs(d, exist_ok=True)
    json.dump(z, open(os.path.join(d, ime + ".json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
